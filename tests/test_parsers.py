from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF

from testgen.ingestion.parsers import UnsupportedDocumentError, parse_document


def write_pdf(path: Path, pages: list[list[str]]) -> Path:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for lines in pages:
        pdf.add_page()
        for line in lines:
            pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))
    return path


def test_markdown_blocks_keep_their_section(tmp_path: Path) -> None:
    md = tmp_path / "srs.md"
    md.write_text(
        "# 1 Introduction\nA library system.\n\n## 3.1 Login\nThe system shall lock accounts.\n",
        encoding="utf-8",
    )
    doc = parse_document(md)
    assert [(b.section, b.text) for b in doc.blocks] == [
        ("1 Introduction", "A library system."),
        ("3.1 Login", "The system shall lock accounts."),
    ]


def test_pdf_blocks_keep_page_and_numbered_section(tmp_path: Path) -> None:
    pdf = write_pdf(
        tmp_path / "srs.pdf",
        [
            ["3.1 User Login", "The system shall lock an account after 5 failed logins."],
            ["The lock shall last 15 minutes.", "3.2 Search", "Users can search by title."],
        ],
    )
    blocks = parse_document(pdf).blocks
    assert [(b.page, b.section) for b in blocks] == [
        (1, "3.1 User Login"),
        (2, "3.1 User Login"),  # the section continues onto the next page
        (2, "3.2 Search"),
    ]
    assert "5 failed logins" in blocks[0].text


def test_numbered_requirement_is_not_mistaken_for_heading(tmp_path: Path) -> None:
    txt = tmp_path / "srs.txt"
    txt.write_text("2 Functional Requirements\n1. The system shall send receipts\n", "utf-8")
    [block] = parse_document(txt).blocks
    assert block.section == "2 Functional Requirements"
    assert block.text == "1. The system shall send receipts"


def test_docx_headings_and_tables_are_read_in_order(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("3.1 Borrowing", level=2)
    document.add_paragraph("Members can borrow books.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "FR-7"
    table.rows[0].cells[1].text = "Loans are limited to 5 books."
    path = tmp_path / "srs.docx"
    document.save(str(path))

    [block] = parse_document(path).blocks
    assert block.section == "3.1 Borrowing"
    assert block.text == "Members can borrow books.\nFR-7 | Loans are limited to 5 books."


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "srs.xlsx"
    path.write_bytes(b"data")
    with pytest.raises(UnsupportedDocumentError, match="Unsupported file type"):
        parse_document(path)


def test_file_without_text_is_rejected(tmp_path: Path) -> None:
    pdf = write_pdf(tmp_path / "scanned.pdf", [[]])
    with pytest.raises(UnsupportedDocumentError, match="No text found"):
        parse_document(pdf)


def test_oversized_file_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    md = tmp_path / "big.md"
    md.write_text("# Title\ntext", encoding="utf-8")
    monkeypatch.setattr("testgen.ingestion.parsers.MAX_FILE_MB", 0)
    with pytest.raises(UnsupportedDocumentError, match="limit is 0 MB"):
        parse_document(md)


def test_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        parse_document(tmp_path / "nope.pdf")
